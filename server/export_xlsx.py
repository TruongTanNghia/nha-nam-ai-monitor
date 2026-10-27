"""Xuất file Excel: dữ liệu gốc cảm biến + bảng dự báo & thực tế + tổng hợp sai số + thời tiết (nộp BTC)."""
import datetime as dt
import io
import time

import config
from forecast_log import HEADERS as FORECAST_HEADERS
from forecast_log import ForecastLog, label
from rf_model import TARGETS
from sensors import DEVICES, SENSORS

RAW_SENSORS = ["t", "h", "GT_AMDAT", "GT_DOMAN", "GT_PH", "ANHSANG"]   # thứ tự cột giống Google Sheet
RAW_DEVICES = ["BOM", "DEN", "QUAT", "SUONG", "MO"]
WEATHER_COLS = [("temp", "Nhiệt độ (°C)"), ("hum", "Độ ẩm (%)"), ("cloud", "Mây (%)"),
                ("rad", "Bức xạ mặt trời (W/m²)"), ("precip", "Mưa (mm)"), ("wind", "Gió (km/h)")]


def _fmt(ts, pattern):
    return dt.datetime.fromtimestamp(ts).strftime(pattern)


def _sheet(wb, title, headers, rows, widths=None):
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title)
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths or [14] * len(headers)):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    head = []
    for h in headers:
        c = WriteOnlyCell(ws, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F6B4F")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        head.append(c)
    ws.append(head)
    for r in rows:
        ws.append(r)
    return ws


def build(engine, days=30):
    """Trả nội dung file .xlsx (bytes)."""
    from openpyxl import Workbook

    now = time.time()
    since = now - days * 86400
    store = engine.store
    wb = Workbook(write_only=True)

    # 1) Giới thiệu
    raw = store.minute_readings(since)
    groups = engine.forecasts.compare(store.forecasts_between(since, now), now)
    done = [g for g in groups if g["actual"]]
    loc = engine.weather.location or {"name": config.WEATHER_NAME, "lat": config.WEATHER_LAT, "lon": config.WEATHER_LON}
    intro = [
        ["Xuất lúc", _fmt(now, "%H:%M %d/%m/%Y")],
        ["Khoảng thời gian", f"{_fmt(since, '%d/%m/%Y')} → {_fmt(now, '%d/%m/%Y')} ({days} ngày)"],
        ["Dữ liệu gốc", f"{len(raw)} phút có số đo. Cảm biến ESP32 đọc qua Firebase 5 giây/lần (khi máy chủ chạy) và "
                        "Google Sheet 1 phút/lần (24/7); mỗi dòng = trung bình các lần đo trong 1 phút. "
                        "Thiết bị: tỉ lệ thời gian bật trong phút (1 = bật cả phút)."],
        ["Dự báo & thực tế", f"{len(groups)} lần dự báo ({len(done)} đã có kết quả). Random Forest dự báo 15 phút/lần "
                             "cho 1, 2, 3 giờ tới; 'thực tế' = trung bình số đo ±7,5 phút quanh giờ được dự báo."],
        ["Tổng hợp sai số", "Sai số tuyệt đối trung bình của Random Forest so với cách 'giữ nguyên giá trị lúc dự báo' "
                            "(mốc chuẩn persistence), trên dữ liệu thật sau khi dự báo; kèm kết quả kiểm chứng "
                            "walk-forward lúc huấn luyện."],
        ["Thời tiết", f"Open-Meteo theo giờ tại {loc.get('name')} ({loc.get('lat')}, {loc.get('lon')})."],
        ["Ghi chú", "Độ ẩm đất ≤ 0,5% coi là cảm biến lỗi/chưa cắm: không dự báo và không chấm điểm. "
                    "Phương pháp chi tiết: docs/PHUONG_PHAP_DU_BAO.md"],
    ]
    _sheet(wb, "Giới thiệu", ["Dữ liệu Nhà Nấm AI", "Nội dung"], intro, [20, 110])

    # 2) Dữ liệu gốc (mỗi phút)
    headers = (["NGÀY", "GIỜ"] + [label(k) for k in RAW_SENSORS] + [DEVICES[d] for d in RAW_DEVICES]
               + ["Chế độ AUTO", "Số lần đo"])
    rows = []
    for r in raw:
        rows.append([_fmt(r["ts"], "%d/%m/%Y"), _fmt(r["ts"], "%H:%M")]
                    + [None if r[k] is None else round(r[k], 2 if k == "GT_PH" else 1) for k in RAW_SENSORS]
                    + [None if r[d] is None else round(r[d], 2) for d in RAW_DEVICES + ["MODE"]] + [r["n"]])
    _sheet(wb, "Dữ liệu gốc", headers, rows, [12, 8] + [14] * (len(headers) - 2))

    # 3) Dự báo & thực tế (giống tab DỰ BÁO trên Google Sheet)
    _sheet(wb, "Dự báo & thực tế", FORECAST_HEADERS, [ForecastLog.row(g) for g in groups],
           [12, 8, 17, 8] + [13] * (len(FORECAST_HEADERS) - 4))

    # 4) Tổng hợp sai số
    rows = [[f"+{s['hours']} giờ", label(s["key"]), s["n"], s["mae"], s["mae_base"], s["skill"],
             "Có" if s["skill"] > 0 else "Chưa"]
            for s in sorted(ForecastLog.score(done), key=lambda s: (TARGETS.index(s["key"]), s["hours"]))]
    rows += [[""] * 7, ["KIỂM CHỨNG WALK-FORWARD LÚC HUẤN LUYỆN (dữ liệu mô hình chưa thấy)"] + [""] * 6]
    rows += [[f"+{e['hours']} giờ", label(e["key"]), e["n"], e["mae"], e["mae_base"], e["skill"],
              "Có" if e["use_rf"] else "Chưa"]
             for e in sorted(engine.rf.info.get("evaluation") or [], key=lambda e: (TARGETS.index(e["key"]), e["hours"]))]
    _sheet(wb, "Tổng hợp sai số",
           ["Dự báo trước", "Chỉ số", "Số lần so", "Sai số TB Random Forest", "Sai số TB giữ nguyên",
            "RF tốt hơn (%)", "RF tốt hơn?"], rows, [13, 26, 11, 16, 16, 13, 11])

    # 5) Thời tiết ngoài trời
    weather = store.weather_between(since, now)
    _sheet(wb, "Thời tiết", ["NGÀY", "GIỜ"] + [n for _, n in WEATHER_COLS],
           [[_fmt(ts, "%d/%m/%Y"), _fmt(ts, "%H:%M")] + [w[k] for k, _ in WEATHER_COLS]
            for ts, w in sorted(weather.items())], [12, 8] + [16] * len(WEATHER_COLS))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
